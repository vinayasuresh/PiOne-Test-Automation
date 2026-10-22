"""Run-owned security review API. No client-supplied arbitrary target URLs."""
import json
from typing import Literal

from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field

from api.services import security_scan_service as service
from api.services.run_context_service import validate_run_id_safe, get_run_context

router = APIRouter(prefix="/runs/{run_id}/security", tags=["Security review"])


class StartReview(BaseModel):
    assessment_mode: Literal["review", "vapt"] = "review"
    authorized: bool = False
    active_opt_in: bool = False
    checks: list[Literal["headers", "cookies", "mixed_content", "disclosure", "api_inventory", "auth_boundary", "error_disclosure", "cors_probe", "reflection_canary"]] = Field(default_factory=lambda: list(service.RULES), min_length=1, max_length=9)
    max_requests: int = Field(default=20, ge=1, le=50)
    max_seconds: int = Field(default=120, ge=10, le=180)


def require_run(run_id):
    try:
        validate_run_id_safe(run_id)
    except ValueError:
        raise HTTPException(400, "Invalid run ID")
    if not get_run_context(run_id):
        raise HTTPException(404, "Run not found")


@router.get("")
def state(run_id: str):
    require_run(run_id)
    return service.snapshot(run_id) or {"run_id": run_id, "status": "not_started"}


@router.post("/start", status_code=202)
def start(run_id: str, request: StartReview):
    require_run(run_id)
    if not request.authorized:
        raise HTTPException(400, "Confirm you have permission to assess this application")
    if set(request.checks) & set(service.ACTIVE_RULES) and not request.active_opt_in:
        raise HTTPException(400, "Active probes require explicit opt-in")
    try:
        return service.start(run_id, list(dict.fromkeys(request.checks)), request.max_requests, request.max_seconds, request.assessment_mode)
    except ValueError as exc:
        raise HTTPException(409, str(exc))


@router.post("/{action}")
def control(run_id: str, action: Literal["pause", "resume", "stop"]):
    require_run(run_id)
    try:
        return service.control(run_id, action)
    except ValueError as exc:
        raise HTTPException(409, str(exc))


@router.get("/export/{format}")
def export(run_id: str, format: Literal["json", "markdown", "excel"]):
    require_run(run_id)
    report = service.snapshot(run_id)
    if not report:
        raise HTTPException(404, "No security findings report for this run")
    if format == "json":
        body, mime, ext = json.dumps(report, indent=2), "application/json", "json"
    elif format == "excel":
        from io import BytesIO
        from openpyxl import Workbook
        wb = Workbook()
        summary = wb.active
        # Keep the historical Coverage sheet name for existing consumers and
        # add a VAPT Summary alias below for the consolidated-report layout.
        summary.title = "Coverage"
        summary.append(["Run", run_id])
        summary.append(["Target", report.get("target", "")])
        summary.append(["Assessment mode", report.get("assessment_mode", "review")])
        summary.append(["Status", report["status"]])
        summary.append(["Requests", report.get("requests", 0)])
        summary.append(["Pages reviewed", report.get("completed_pages", 0)])
        summary.append(["Findings", len(report.get("findings", []))])
        for severity, count in sorted((report.get("summary") or {}).get("by_severity", {}).items()):
            summary.append([f"Severity: {severity}", count])
        vapt_summary = wb.create_sheet("VAPT Summary")
        for row in summary.iter_rows(values_only=True):
            vapt_summary.append(list(row))
        for warning in report["warnings"] + report["not_tested"]:
            summary.append(["Limitation", warning])
        findings = wb.create_sheet("Master Findings")
        fields = ["finding_id", "source_finding_id", "category", "title", "description", "severity", "cvss", "cve", "owasp_mapping", "affected_components", "test_scenario", "expected_result", "actual_result", "evidence", "business_impact", "remediation", "status", "source_file", "source_sheet", "platform", "module", "id", "url", "method", "confidence", "verification"]
        findings.append(fields)
        for finding in report["findings"]:
            findings.append([str(finding.get(key, "")) for key in fields])
        endpoints = wb.create_sheet("Endpoint Inventory")
        endpoints.append(["URL", "Method", "Status", "Source"])
        for endpoint in report.get("endpoints", []):
            endpoints.append([endpoint.get("url", ""), endpoint.get("method", "GET"), endpoint.get("status", ""), endpoint.get("source", "route")])
        gaps = wb.create_sheet("Coverage Gaps")
        gaps.append(["Coverage gap"])
        for warning in report.get("not_tested", []):
            gaps.append([warning])
        # Force strings to remain data, including formula-looking target paths.
        for sheet in wb:
            for row in sheet:
                for cell in row:
                    if cell.value is not None:
                        cell.value = str(cell.value)
                    cell.data_type = "s"
            sheet.freeze_panes = "A2"
        output = BytesIO()
        wb.save(output)
        body, mime, ext = output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    else:
        # Escape Markdown/HTML from target-controlled fields in downloadable reports.
        def clean(value):
            import html
            return html.escape(str(value)).replace("`", "\\`").replace("[", "\\[").replace("]", "\\]")
        lines = ["# PiOne Testmate VAPT assessment", f"Run: {run_id}", f"Target: {report.get('target', '')}", f"Mode: {report.get('assessment_mode', 'review')}", f"Status: {report['status']}",
                 "Authorized, bounded assessment. No destructive actions, credential attacks, brute force, denial of service, or mutation requests are performed.", "## Coverage limitations"]
        lines += ["- " + clean(x) for x in report["warnings"]]
        lines += ["Not tested: " + ", ".join(report["not_tested"])]
        for f in report["findings"]:
            lines += ["", "## " + clean(f["title"])]
            lines += [f"{key}: {clean(f.get(key, ''))}" for key in ("finding_id", "category", "title", "url", "severity", "cvss", "owasp_mapping", "status", "confidence", "verification", "evidence", "business_impact", "remediation")]
        body, mime, ext = "\n\n".join(lines), "text/markdown", "md"
    return Response(body, media_type=mime, headers={"Content-Disposition": f'attachment; filename="security_{run_id}.{ext}"', "Cache-Control": "no-store"})
